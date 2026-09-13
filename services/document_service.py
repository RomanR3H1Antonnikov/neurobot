"""
Этап 2: работа с документами.
Парсинг файла → передача текста в ИИ вместе с заданием пользователя.
"""
import io
from providers.base import TaskType
from providers.router import get_provider, get_task_rate_limit
from db.queries import (
    get_or_create_user, deduct_credits, check_and_increment_rate_limit,
)
from services.media_service import InsufficientCreditsError, RateLimitError

MAX_CHARS = 40_000  # обрезаем документ, чтобы не выйти за лимит токенов модели


def extract_text(filename: str, file_bytes: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        return _extract_pdf(file_bytes)
    elif ext in ("docx", "doc"):
        return _extract_docx(file_bytes)
    elif ext in ("xlsx", "xls"):
        return _extract_xlsx(file_bytes)
    else:
        return file_bytes.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
    return "\n".join(text_parts)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"[Лист: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            rows.append("\t".join(str(c) if c is not None else "" for c in row))
    return "\n".join(rows)


async def process_document(
    telegram_id: int, username: str, filename: str, file_bytes: bytes, task_text: str
) -> str:
    user = await get_or_create_user(telegram_id, username)
    user_id = user["id"]

    rate_limit = get_task_rate_limit(TaskType.DOCUMENT)
    ok = await check_and_increment_rate_limit(user_id, TaskType.DOCUMENT.value, rate_limit)
    if not ok:
        raise RateLimitError(f"Превышен лимит запросов ({rate_limit} в час)")

    provider, model_cfg = get_provider(TaskType.DOCUMENT)
    cost = model_cfg["cost_credits"]

    if user["balance"] < cost:
        raise InsufficientCreditsError(
            f"Недостаточно средств. Нужно: {cost} ₽, у вас: {user['balance']} ₽"
        )

    doc_text = extract_text(filename, file_bytes)
    if len(doc_text) > MAX_CHARS:
        doc_text = doc_text[:MAX_CHARS] + "\n...[документ обрезан]"

    messages = [{
        "role": "user",
        "content": (
            f"Вот содержимое документа «{filename}»:\n\n{doc_text}\n\n"
            f"Задание: {task_text}"
        ),
    }]

    result = await provider.chat(messages, model=model_cfg["model_id"])
    await deduct_credits(user_id, cost, TaskType.DOCUMENT.value)
    return result.text
